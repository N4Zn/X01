"""Xuất Excel kết quả (cần openpyxl)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from edux_lib import LEVEL_NAMES, level


def build(cfg, out, stats, excluded, not_roster, path, period, by_day=None):
    f = cfg['formula']
    HP = cfg['hp_order']
    lv = lambda p: LEVEL_NAMES[level(f, p)]
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tổng hợp'
    hdr = ['ID', 'Tên đầy đủ', 'Tên thường gọi'] + HP + ['Kết quả tổng (%s)' % cfg['mon'], 'Mức', 'Số câu tính điểm', 'Ghi chú']
    ws.append(['KẾT QUẢ ĐÁNH GIÁ %s – %s – Điểm học phần = %g × tỉ lệ đúng × hệ số tốc độ (%g nếu TB ≤%gs; 1,0 nếu ≥%gs; tuyến tính ở giữa; tối đa 100). Tổng = TB các học phần có chơi. Ô vàng = học phần ít mẫu (<%d câu). Nguồn: Google Sheet log online.'
               % (cfg['mon'].upper(), period, f['base'], f['kmax'], f['tfast'], f['tslow'], f['low_sample'])])
    ws.append(hdr)
    thin = Side(style='thin', color='BBBBBB')
    yel = PatternFill('solid', fgColor='FFF2CC')
    head = PatternFill('solid', fgColor='1F3864')
    for d in [o for o in out if 'toan' in o]:
        r = [d['code'], d['name'], d['alias']]
        for h in HP:
            r.append(round(d['hp'][h]['pt'], 1) if h in d['hp'] else None)
        notes = []
        r += [round(d['toan'], 1), lv(d['toan']), d['n']]
        if d['n'] < f['few_total']:
            notes.append('ÍT DỮ LIỆU: chỉ %d câu tính điểm' % d['n'])
        low = [h for h in HP if h in d['hp'] and d['hp'][h]['low']]
        if low:
            notes.append('Học phần ít mẫu: ' + ', '.join(low))
        miss = [h for h in HP if h not in d['hp']]
        if miss:
            notes.append('CHƯA CÓ ĐIỂM: ' + ', '.join(miss))
        r.append('; '.join(notes))
        ws.append(r)
        i = ws.max_row
        for j, h in enumerate(HP):
            if h in d['hp'] and d['hp'][h]['low']:
                ws.cell(i, 4 + j).fill = yel
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(hdr))
    ws['A1'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[1].height = 48
    for c in ws[2]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        row[3 + len(HP)].font = Font(bold=True)
    for col, w in zip('ABCDEFGHIJKL', [13, 24, 14, 13, 9, 9, 9, 10, 16, 14, 12, 52]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'D3'

    w2 = wb.create_sheet('Chi tiết học phần')
    w2.append(['ID', 'Tên thường gọi', 'Học phần', 'Đúng', 'Số câu', '% đúng', 'TB giây (câu đúng)', 'Hệ số tốc độ', 'Điểm', 'Ít mẫu'])
    for d in out:
        for h in HP:
            if h in d['hp']:
                v = d['hp'][h]
                w2.append([d['code'], d['alias'], h, v['ok'], v['n'], round(v['acc'], 1), round(v['t'], 1) if v['t'] else None,
                           round(v['k'], 2), round(v['pt'], 1), 'x' if v['low'] else ''])
    for c in w2[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head
    for col, w in zip('ABCDEFGHIJ', [13, 14, 14, 7, 8, 8, 18, 13, 8, 8]):
        w2.column_dimensions[col].width = w

    w3 = wb.create_sheet('Loại trừ & chất lượng log')
    w3.append(['Thống kê nạp log (Google Sheet)'])
    labels = [('rows', 'Tổng dòng trên Sheet'),
              ('duplicates', 'Dòng trùng (gửi lại khi mạng chập chờn)'),
              ('bad_clock_or_test', 'Dòng đồng hồ máy sai năm / TEST'),
              ('no_result', 'round_end thiếu kết quả/thời gian'),
              ('unassigned', 'Round chưa gán được học sinh (không có nhận diện trong %ds)' % cfg['recog_max_gap_sec']),
              ('not_in_roster', 'Round của tên ngoài danh sách lớp (thử nghiệm/giáo viên)'),
              ('rounds', 'Round gán được cho học sinh'),
              ('via_recognition', '… trong đó gán nhờ nhận diện gần nhất')]
    for k, lab in labels:
        w3.append([lab, stats.get(k, 0)])
    w3.append([])
    w3.append(['Round bị loại khi tính điểm (kỳ %s)' % period])
    for k, v in sorted(excluded.items(), key=lambda x: -x[1]):
        w3.append([k, v])
    if not_roster:
        w3.append([])
        w3.append(['Tên ngoài danh sách lớp (số round)'])
        for k, v in not_roster.most_common(20):
            w3.append([k or '(trống)', v])
    w3['A1'].font = Font(bold=True)
    w3.column_dimensions['A'].width = 70
    w3.column_dimensions['B'].width = 14

    if by_day:
        wd = wb.create_sheet('Theo ngày', 1)
        days = sorted(by_day)
        wd.append(['ID', 'Tên thường gọi'] + sum([['%s: điểm' % d[5:].replace('-', '/'), 'số câu'] for d in days], []) + ['Cả kỳ: điểm', 'số câu'])
        by = {d: {o['code']: o for o in by_day[d]} for d in days}
        for o in out:
            r = [o['code'], o['alias']]
            for d in days:
                x = by[d].get(o['code'])
                r += [round(x['toan'], 1) if x and 'toan' in x else None, x['n'] if x else 0]
            r += [round(o['toan'], 1) if 'toan' in o else None, o['n']]
            wd.append(r)
        for c in wd[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = head
            c.alignment = Alignment(horizontal='center', wrap_text=True)
        wd.column_dimensions['A'].width = 13
        wd.column_dimensions['B'].width = 16

    w4 = wb.create_sheet('Chưa có điểm')
    w4.append(['ID', 'Tên đầy đủ', 'Tên thường gọi', 'Ghi chú'])
    for d in [o for o in out if 'toan' not in o]:
        w4.append([d['code'], d['name'], d['alias'], 'Chưa có câu nào ở game tính điểm trong kỳ này.'])
    for c in w4[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = head
    for col, w in zip('ABCD', [13, 24, 14, 70]):
        w4.column_dimensions[col].width = w
    wb.save(path)
